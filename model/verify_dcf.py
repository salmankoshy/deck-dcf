# -*- coding: utf-8 -*-
"""Verify model/deck_dcf.xlsx: (1) independent Python DCF from the CSVs, (2) recalc the
workbook in Excel via COM (LibreOffice not installed) and compare, (3) scan WACC/DCF/Sensitivity
for typed numbers. Run: py -3.13 model/verify_dcf.py"""
import csv, os, sys, shutil, tempfile
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "model", "deck_dcf.xlsx")

def read_csv(p):
    with open(p, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))
A = {r["key"]: r["value"] for r in read_csv(os.path.join(ROOT, "inputs", "assumptions.csv"))}
H = {(r["item"], r["period"]): r for r in read_csv(os.path.join(ROOT, "inputs", "historicals.csv"))}
fs = lambda item, p="FY2026": float(H[(item, p)]["fs_value"])
f = lambda k: float(A[k])

# ---------------------------------------------------------------- 1. Python DCF
rf, beta, erp = f("rf"), f("beta"), f("erp")
kd, tax = f("cost_debt_pretax"), f("tax_rate")
price, shares, cash, debt = f("price"), f("shares_diluted_current"), f("cash"), f("debt")
tg, exit_mult = f("terminal_growth"), f("exit_multiple_ev_ebitda")
g = [f(f"rev_growth_fy{y}") for y in (27, 28, 29, 30, 31)]
m, da_pct, capex_pct, nwc_pct = f("ebit_margin"), f("da_pct_rev"), f("capex_pct_rev"), f("nwc_pct_rev")

ke = rf + beta * erp
mcap = price * shares / 1000
we, wd = mcap / (mcap + debt), debt / (mcap + debt)
wacc = ke * we + kd * (1 - tax) * wd

rev26 = fs("Net sales")
nwc26 = fs("Trade accounts receivable net") + fs("Inventories") - fs("Trade accounts payable")
years = ["FY2027E", "FY2028E", "FY2029E", "FY2030E", "FY2031E"]
rev, ebit, nopat, da, capex, nwc, dnwc, fcff, ebitda, t, df, pv = ([] for _ in range(12))
prev_rev, prev_nwc = rev26, nwc26
for i, gi in enumerate(g):
    r_ = prev_rev * (1 + gi); rev.append(r_)
    e = r_ * m; ebit.append(e)
    nopat.append(e * (1 - tax))
    da.append(da_pct * r_); capex.append(capex_pct * r_)
    n = nwc_pct * r_; nwc.append(n); dnwc.append(n - prev_nwc)
    fcff.append(nopat[-1] + da[-1] - capex[-1] - dnwc[-1])
    ebitda.append(e + da[-1])
    t.append(i + 0.5); df.append(1 / (1 + wacc) ** t[-1]); pv.append(fcff[-1] * df[-1])
    prev_rev, prev_nwc = r_, n
sum_pv = sum(pv)
tv_p = fcff[-1] * (1 + tg) / (wacc - tg); pv_tv_p = tv_p / (1 + wacc) ** 4.5
tv_x = ebitda[-1] * exit_mult;             pv_tv_x = tv_x / (1 + wacc) ** 5.0
ev_p, ev_x = sum_pv + pv_tv_p, sum_pv + pv_tv_x
eq_p, eq_x = ev_p + cash - debt, ev_x + cash - debt
px_p, px_x = eq_p * 1000 / shares, eq_x * 1000 / shares

def px_margin(m_):
    fc = [r_ * m_ * (1 - tax) + d_ - c_ - n_ for r_, d_, c_, n_ in zip(rev, da, capex, dnwc)]
    pvs = sum(x / (1 + wacc) ** tt for x, tt in zip(fc, t))
    return (pvs + fc[-1] * (1 + tg) / (wacc - tg) / (1 + wacc) ** t[-1] + cash - debt) / shares * 1000
m_lo, m_hi = f("implied_margin_probe_lo"), f("implied_margin_probe_hi")
implied_m = m_lo + (price - px_margin(m_lo)) * (m_hi - m_lo) / (px_margin(m_hi) - px_margin(m_lo))
tv_x_mid = tv_x / (1 + wacc) ** 0.5          # exit TV pulled back to the mid-year point (t = 4.5)
impx_timed = tv_p * (1 + wacc) ** 0.5 / ebitda[-1]   # perpetuity TV pushed to the end-year point (t = 5.0)
impg_timed = (tv_x_mid * wacc - fcff[-1]) / (tv_x_mid + fcff[-1])
checks = {
    "TV % of EV (perp) 60-80%": (pv_tv_p / ev_p, 0.60 <= pv_tv_p / ev_p <= 0.80),
    "Implied exit multiple, timed to exit convention, 8-12x": (impx_timed, 8 <= impx_timed <= 12),
    "Implied g from exit TV, timed to exit convention, 1.5-3.5%": (impg_timed, 0.015 <= impg_timed <= 0.035),
    "Methods within +/-20%": (abs(px_p / px_x - 1), abs(px_p / px_x - 1) <= 0.20),
    "Terminal FCFF margin 12-18%": (fcff[-1] / rev[-1], 0.12 <= fcff[-1] / rev[-1] <= 0.18),
    "Y1 DF = 1/(1+WACC)^0.5": (df[0] - 1 / (1 + wacc) ** 0.5, abs(df[0] - 1 / (1 + wacc) ** 0.5) < 1e-12),
    "Price at market-implied margin equals price": (px_margin(implied_m), abs(px_margin(implied_m) - price) < 1e-6),
}

K = lambda v: f"{v:,.0f}"
print("=" * 78); print("1. INDEPENDENT PYTHON DCF (from inputs/*.csv)"); print("=" * 78)
print(f"Ke = {rf} + {beta}*{erp} = {ke:.6f}   We={we:.4f} Wd={wd:.4f}   WACC = {wacc:.6f}")
print(f"Base FY2026A: revenue {K(rev26)}  NWC (AR+Inv-AP) {K(nwc26)}")
print(f"\n{'':10}" + "".join(f"{y:>14}" for y in years))
for lab, arr in [("growth", g), ("revenue", rev), ("EBIT", ebit), ("NOPAT", nopat), ("D&A", da), ("capex", capex),
                 ("NWC", nwc), ("dNWC", dnwc), ("FCFF", fcff), ("EBITDA", ebitda), ("t", t), ("DF", df), ("PV FCFF", pv)]:
    if lab in ("growth",): row = "".join(f"{v:>14.3%}" for v in arr)
    elif lab in ("t",): row = "".join(f"{v:>14.1f}" for v in arr)
    elif lab in ("DF",): row = "".join(f"{v:>14.5f}" for v in arr)
    else: row = "".join(f"{K(v):>14}" for v in arr)
    print(f"{lab:10}{row}")
print(f"\n{'':32}{'Perpetuity':>16}{'Exit mult':>16}")
for lab, a, b in [("TV (undiscounted)", tv_p, tv_x), ("PV of TV", pv_tv_p, pv_tv_x), ("Sum PV FCFF", sum_pv, sum_pv),
                  ("EV", ev_p, ev_x), ("+cash -debt -> equity", eq_p, eq_x)]:
    print(f"{lab:32}{K(a):>16}{K(b):>16}")
print(f"{'Price / share ($)':32}{px_p:>16.4f}{px_x:>16.4f}")
print(f"{'Upside vs ' + str(price):32}{px_p / price - 1:>16.2%}{px_x / price - 1:>16.2%}")
print("\nChecks (Python):")
for k, (v, ok) in checks.items():
    vs = f"{v:.4%}" if abs(v) < 1 else f"{v:,.4f}"
    print(f"  {'PASS' if ok else 'FAIL'}  {k:45} {vs}")

# ---- independent sensitivity grids (same maths the Sensitivity tab formulas implement)
s_w, s_g, s_m = f("sens_wacc_step"), f("sens_g_step"), f("sens_mult_step")
KS = [-2, -1, 0, 1, 2]
w_axis = [wacc + k * s_w for k in KS]
g_axis = [tg + k * s_g for k in KS]
m_axis = [exit_mult + k * s_m for k in KS]
def px_perp(w, g_):
    pvs = sum(fc / (1 + w) ** tt for fc, tt in zip(fcff, t))
    return (pvs + fcff[-1] * (1 + g_) / (w - g_) / (1 + w) ** t[-1] + cash - debt) / shares * 1000
def px_exit(w, m_):
    pvs = sum(fc / (1 + w) ** tt for fc, tt in zip(fcff, t))
    return (pvs + ebitda[-1] * m_ / (1 + w) ** 5.0 + cash - debt) / shares * 1000
py_grid1 = [[px_perp(w, g_) for g_ in g_axis] for w in w_axis]
py_grid2 = [[px_exit(w, m_) for m_ in m_axis] for w in w_axis]
# operating beta from the regression, ex-IFRS peer median, value at the current multiple
beta5_raw = float(next(r_ for r_ in read_csv(os.path.join(ROOT, "inputs", "beta.csv")) if r_["window"].startswith("5Y"))["raw_beta"])
b_op = beta5_raw / (1 - cash / mcap)
ke_op = rf + b_op * erp
wacc_op = ke_op * we + kd * (1 - tax) * wd
px_op = px_perp(wacc_op, tg)
# ROIC FY26 (Historical), implied RONIC and the RONIC-based terminal rows (DCF)
eq26, cash26 = fs("Total stockholders' equity"), fs("Cash and cash equivalents")
oll26 = fs("Operating lease liabilities current") + fs("Long-term operating lease liabilities")
ebit26 = fs("Income from operations")
roic_py = ebit26 * (1 - tax) / (eq26 - cash26 + oll26)
roic_ex_py = ebit26 * (1 - tax) / (eq26 - cash26)
ronic_in = f("ronic")
ronic_implied_py = tg / ((nopat[-1] - fcff[-1]) / nopat[-1])
px_ronic_py = (sum_pv + nopat[-1] * (1 + tg) * (1 - tg / ronic_in) / (wacc - tg) / (1 + wacc) ** t[-1] + cash - debt) / shares * 1000
N_, T_ = nopat[-1], tv_x_mid
qa_, qb_, qc_ = -N_ / ronic_in, N_ * (1 - 1 / ronic_in) + T_, N_ - T_ * wacc
g_ronic_py = (-qb_ + (qb_ ** 2 - 4 * qa_ * qc_) ** 0.5) / (2 * qa_)
print(f"ROIC FY26 {roic_py:.4%} (ex leases {roic_ex_py:.4%}); implied RONIC {ronic_implied_py:.4%}; price at RONIC {ronic_in:.0%}: ${px_ronic_py:.2f}; g from exit at that RONIC {g_ronic_py:.4%}")
print(f"\nmarket-implied EBIT margin {implied_m:.4%} (price at it {px_margin(implied_m):.4f}); operating beta {b_op:.4f} -> Ke {ke_op:.4%} -> ${px_op:.2f}")

def show_grid(title, rows_axis, cols_axis, body, col_fmt, corner="WACC \\ "):
    print(f"\n{title}")
    print(f"{corner:>10}" + "".join(f"{col_fmt(c):>12}" for c in cols_axis))
    for w, row in zip(rows_axis, body):
        print(f"{w:>10.2%}" + "".join(f"{v:>12.2f}" for v in row))


# ---- LTM EBITDA to Jun 30 2026, comps implied prices, football field (same maths as the Comps / Historical tabs)
import numpy as np
from openpyxl import load_workbook as _lw
_wbf = _lw(XLSX)   # formulas mode: read the typed (blue) Comps inputs so nothing is hardcoded here
ltm_ebit_py = fs("Income from operations", "FY2026") + fs("Income from operations", "Q1FY27") - fs("Income from operations", "Q1FY26")
ltm_da_py = (fs("Depreciation amortization and accretion", "FY2026") + fs("Depreciation amortization and accretion", "Q1FY27")
             - fs("Depreciation amortization and accretion", "Q1FY26"))
ltm_ebitda_py = ltm_ebit_py + ltm_da_py
csv_ltm_ebitda = fs("Income from operations", "LTM_Jun26") + fs("Depreciation amortization and accretion", "LTM_Jun26")
_sc = _wbf["Comps"]
peer_mults = [(_sc[f"A{rr}"].value, float(_sc[f"B{rr}"].value)) for rr in range(4, 40)
              if _sc[f"B{rr}"].data_type == "n" and isinstance(_sc[f"A{rr}"].value, str) and _sc[f"A{rr}"].value.isupper()]
_m = np.array([v for _, v in peer_mults])
q1_py, med_py, q3_py = (float(np.percentile(_m, p)) for p in (25, 50, 75))   # linear interpolation = Excel QUARTILE / MEDIAN
px_comps_py = {k: (q * ltm_ebitda_py + cash - debt) * 1000 / shares for k, q in (("q1", q1_py), ("med", med_py), ("q3", q3_py))}
deck_ev_ebitda_py = (price * shares / 1000 + debt - cash) / ltm_ebitda_py
med_ex_ifrs_py = float(np.median([v for k, v in peer_mults if k in ("CROX", "NKE", "WWW", "SHOO")]))
px_cur_mult_py = px_exit(wacc, deck_ev_ebitda_py)
_inner = lambda grid: [v for row in grid[1:4] for v in row[1:4]]
_mc = {r_["metric"]: float(r_["value"]) for r_ in read_csv(os.path.join(ROOT, "mc", "mc_summary.csv"))} if os.path.exists(os.path.join(ROOT, "mc", "mc_summary.csv")) else {}
_ff_typed = {}
for rr in range(1, _sc.max_row + 1):
    lab = _sc[f"A{rr}"].value
    if lab in ("Analyst price targets", "52-week range"):
        _ff_typed[lab] = (float(_sc[f"B{rr}"].value), float(_sc[f"C{rr}"].value))
ff_py = {
    "dcf_perp": (min(_inner(py_grid1)), max(_inner(py_grid1))),
    "dcf_exit": (min(_inner(py_grid2)), max(_inner(py_grid2))),
    "mc": (_mc.get("P10", float("nan")), _mc.get("P90", float("nan"))),
    "comps": (px_comps_py["q1"], px_comps_py["q3"]),
    "pt": _ff_typed.get("Analyst price targets", (float("nan"),) * 2),
    "w52": _ff_typed.get("52-week range", (float("nan"),) * 2),
}
print("\n" + "=" * 78); print("1b. LTM EBITDA / COMPS / FOOTBALL FIELD (Python from inputs/*.csv + typed Comps inputs)"); print("=" * 78)
print(f"LTM income from operations = {fs('Income from operations', 'FY2026'):,.0f} + {fs('Income from operations', 'Q1FY27'):,.0f} - {fs('Income from operations', 'Q1FY26'):,.0f} = {ltm_ebit_py:,.0f}")
print(f"LTM D&A and accretion      = {fs('Depreciation amortization and accretion', 'FY2026'):,.0f} + {fs('Depreciation amortization and accretion', 'Q1FY27'):,.0f} - {fs('Depreciation amortization and accretion', 'Q1FY26'):,.0f} = {ltm_da_py:,.0f}")
print(f"LTM EBITDA = {ltm_ebitda_py:,.0f}   (csv LTM_Jun26 rows sum to {csv_ltm_ebitda:,.0f}: {'match' if csv_ltm_ebitda == ltm_ebitda_py else 'MISMATCH'})")
print("peer multiples (typed on Comps tab):", ", ".join(f"{k} {v:.2f}x" for k, v in peer_mults))
print(f"peer 25th / median / 75th: {q1_py:.4f}x / {med_py:.4f}x / {q3_py:.4f}x")
print(f"implied price per share: 25th ${px_comps_py['q1']:.4f}   median ${px_comps_py['med']:.4f}   75th ${px_comps_py['q3']:.4f}")
print(f"DECK current EV / LTM EBITDA: {deck_ev_ebitda_py:.4f}x")

# ---------------------------------------------------------------- 2. Excel recalc via COM
print("\n" + "=" * 78); print("2. WORKBOOK RECALC"); print("=" * 78)
print("LibreOffice: not installed. Microsoft Excel is installed -> recalculating via Excel COM instead.")
xl_vals = {}
errors = []
try:
    import win32com.client as w32
    import pythoncom
    pythoncom.CoInitialize()
    tmpdir = tempfile.mkdtemp()
    tmp = os.path.join(tmpdir, "deck_dcf_recalc.xlsx")
    shutil.copy(XLSX, tmp)
    xl = w32.DispatchEx("Excel.Application")
    xl.Visible = False; xl.DisplayAlerts = False
    wb = xl.Workbooks.Open(tmp)
    xl.CalculateFullRebuild()
    EXTRA = ["ltm_ebit", "ltm_da", "ltm_ebitda", "comps_q1", "comps_med", "comps_q3", "px_comps_q1", "px_comps_med", "px_comps_q3",
             "deck_ev_ebitda", "ff_base_dcf", "ff_price", "mc_recon_check", "mc_price_check", "implied_margin", "px_at_implied_margin",
             "beta_operating", "ke_operating", "wacc_operating", "px_at_ke_operating", "comps_med_ex_ifrs", "px_exit_at_current_mult",
             "integrity_summary", "cross_check_flags", "implied_ev_ltm_ebitda", "implied_ev_fy27_ebitda", "roic_fy26", "roic_fy26_ex_leases", "ronic_implied", "px_at_ronic", "g_exit_at_ronic"] + [f"ff_{k}_{s}" for k in ("dcf_perp", "dcf_exit", "mc", "comps", "pt", "w52") for s in ("lo", "hi")]
    for nm in ["wacc_calc", "ke_calc", "price_perp", "price_exit", "upside_perp", "upside_exit", "ev_perp", "ev_exit"] + EXTRA:
        xl_vals[nm] = wb.Names(nm).RefersToRange.Value
    # every check result + ALL CHECKS on DCF
    dcf = wb.Worksheets("DCF")
    used = dcf.UsedRange
    xl_checks = []
    for row in range(1, used.Rows.Count + 1):
        d = dcf.Cells(row, 4).Value
        if d in ("PASS", "FAIL", "in range", "outside range", "context"):
            xl_checks.append((dcf.Cells(row, 1).Value, dcf.Cells(row, 2).Value, d))
    # scan every sheet for error values
    for sh in wb.Worksheets:
        ur = sh.UsedRange
        for row in range(1, ur.Rows.Count + 1):
            for col in range(1, ur.Columns.Count + 1):
                c = ur.Cells(row, col)
                if c.HasFormula:
                    txt = str(c.Text)
                    if txt.startswith("#"):
                        errors.append(f"{sh.Name}!{c.Address(False, False)} -> {txt}  [{c.Formula}]")
    ncells_formula = sum(1 for sh in wb.Worksheets for c in sh.UsedRange if c.HasFormula)
    # sensitivity grids: header rows are where column B reads "WACC ↓ ..."
    sens = wb.Worksheets("Sensitivity")
    su = sens.UsedRange
    xl_grids = []
    for row in range(1, su.Rows.Count + 1):
        v = sens.Cells(row, 2).Value
        if isinstance(v, str) and v.startswith("WACC"):
            cols = [sens.Cells(row, c).Value for c in range(3, 8)]
            rws, body = [], []
            for i in range(1, 6):
                rws.append(sens.Cells(row + i, 2).Value)
                body.append([sens.Cells(row + i, c).Value for c in range(3, 8)])
            xl_grids.append((sens.Cells(row - 2, 1).Value, rws, cols, body))
    xl_sens_checks = {}
    for nm in ("sens_check_perp", "sens_check_exit"):
        rr_ = wb.Names(nm).RefersToRange
        # PASS/FAIL sits in column C; the center − price difference is the cell to its left (column B)
        xl_sens_checks[nm] = (rr_.Value, sens.Cells(rr_.Row, rr_.Column - 1).Value)
    wb.Close(SaveChanges=False); xl.Quit()
    print(f"formula cells in workbook: {ncells_formula}")
    print(f"error cells (#REF!/#DIV/0!/#NAME?/...): {len(errors)}")
    for e in errors: print("   ", e)
    cmp = [("WACC", xl_vals["wacc_calc"], wacc), ("Ke", xl_vals["ke_calc"], ke),
           ("price_perp", xl_vals["price_perp"], px_p), ("price_exit", xl_vals["price_exit"], px_x),
           ("EV perp", xl_vals["ev_perp"], ev_p), ("EV exit", xl_vals["ev_exit"], ev_x),
           ("upside_perp", xl_vals["upside_perp"], px_p / price - 1), ("upside_exit", xl_vals["upside_exit"], px_x / price - 1)]
    print(f"\n{'cell':14}{'Excel':>18}{'Python':>18}{'diff':>14}")
    allok = True
    for lab, a, b in cmp:
        d = a - b; ok = abs(d) < 1e-6 * max(1, abs(b))
        allok &= ok
        print(f"{lab:14}{a:>18.6f}{b:>18.6f}{d:>14.2e}  {'OK' if ok else 'MISMATCH'}")
    print("Excel == Python:", "YES" if allok else "NO")
    print("\nChecks (Excel):")
    for lab, v, res in xl_checks:
        print(f"  {res}  {lab}  = {v}")

    print("\n" + "=" * 78); print("SENSITIVITY GRIDS (values as recalculated by Excel)"); print("=" * 78)
    fmts = [lambda c: f"{c:.2%}", lambda c: f"{c:.1f}x"]
    pyg = [py_grid1, py_grid2]
    for (title, rws, cols, body), cf, pg in zip(xl_grids, fmts, pyg):
        show_grid(title, rws, cols, body, cf)
        maxdiff = max(abs(a - b) for ra, rb in zip(body, pg) for a, b in zip(ra, rb))
        print(f"   max |Excel − Python| over 25 cells: {maxdiff:.2e}")
    print("\nCenter checks (Excel):")
    for nm, (res, diff) in xl_sens_checks.items():
        print(f"  {res}  {nm}: center − DCF price = {diff:.2e}")

    print("\n" + "=" * 78); print("2b. LTM / COMPS / FOOTBALL FIELD (Excel vs Python)"); print("=" * 78)
    cmp2 = [("ltm_ebit", ltm_ebit_py), ("ltm_da", ltm_da_py), ("ltm_ebitda", ltm_ebitda_py),
            ("comps_q1", q1_py), ("comps_med", med_py), ("comps_q3", q3_py),
            ("px_comps_q1", px_comps_py["q1"]), ("px_comps_med", px_comps_py["med"]), ("px_comps_q3", px_comps_py["q3"]),
            ("deck_ev_ebitda", deck_ev_ebitda_py), ("ff_base_dcf", px_p), ("ff_price", price)]
    for k, (lo, hi) in ff_py.items():
        cmp2 += [(f"ff_{k}_lo", lo), (f"ff_{k}_hi", hi)]
    cmp2 += [("implied_margin", implied_m), ("px_at_implied_margin", price), ("beta_operating", b_op), ("ke_operating", ke_op),
             ("wacc_operating", wacc_op), ("px_at_ke_operating", px_op), ("comps_med_ex_ifrs", med_ex_ifrs_py), ("px_exit_at_current_mult", px_cur_mult_py),
             ("roic_fy26", roic_py), ("roic_fy26_ex_leases", roic_ex_py), ("ronic_implied", ronic_implied_py), ("px_at_ronic", px_ronic_py), ("g_exit_at_ronic", g_ronic_py),
             ("implied_ev_ltm_ebitda", ev_p / ltm_ebitda_py), ("implied_ev_fy27_ebitda", ev_p / ebitda[0])]
    print(f"{'name':16}{'Excel':>18}{'Python':>18}{'diff':>14}")
    allok2 = True
    for nm, b in cmp2:
        a = xl_vals[nm]; d = a - b; ok = abs(d) < 1e-6 * max(1, abs(b))
        allok2 &= ok
        print(f"{nm:16}{a:>18.6f}{b:>18.6f}{d:>14.2e}  {'OK' if ok else 'MISMATCH'}")
    print("Excel == Python (LTM / comps / football):", "YES" if allok2 else "NO")
    print("mc_recon_check:", xl_vals["mc_recon_check"], "| mc_price_check:", xl_vals["mc_price_check"])
    print(f"integrity: {xl_vals['integrity_summary']} | cross-check flags: {xl_vals['cross_check_flags']}")
except Exception as ex:
    print("Excel COM recalc failed:", repr(ex))

# ---------------------------------------------------------------- 3. typed-number scan
print("\n" + "=" * 78); print("3. TYPED-NUMBER SCAN: WACC / DCF / Sensitivity"); print("=" * 78)
from openpyxl import load_workbook
wbo = load_workbook(XLSX)  # formulas as strings
hits = []
for sn in ("WACC", "DCF", "Sensitivity"):
    for row in wbo[sn].iter_rows():
        for c in row:
            if c.value is None: continue
            if c.data_type == "n":          # numeric constant (formulas are 'f', text 's')
                hits.append(f"{sn}!{c.coordinate} = {c.value!r}")
print(f"cells holding a typed numeric constant: {len(hits)}")
for h in hits: print("   ", h)
# also report formula cells that are just a bare constant (e.g. '=0.08')
bare = []
for sn in ("WACC", "DCF", "Sensitivity"):
    for row in wbo[sn].iter_rows():
        for c in row:
            if c.data_type == "f":
                body = str(c.value)[1:].strip()
                try:
                    float(body); bare.append(f"{sn}!{c.coordinate} {c.value}")
                except ValueError: pass
print(f"formula cells that are a bare constant (e.g. '=0.08' axis anchors): {len(bare)}")
for b in bare: print("   ", b)
