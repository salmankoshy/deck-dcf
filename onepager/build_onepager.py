# -*- coding: utf-8 -*-
"""Build onepager/DECK_pitch.pdf (one US Letter page) from HTML via headless Edge/Chrome --print-to-pdf.
Every model number is read from model/deck_dcf.xlsx (openpyxl data_only=True) through defined names or
label-located cells; filing facts are typed here with their source. Writes onepager/number_map.csv and
onepager/preview.png (page 1 rasterized). Run: py -3.13 onepager/build_onepager.py
"""
import csv, os, sys, subprocess, shutil
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ONE = os.path.join(ROOT, "onepager")
XLSX = os.path.join(ROOT, "model", "deck_dcf.xlsx")
HTML = os.path.join(ONE, "DECK_pitch.html")
PDF = os.path.join(ONE, "DECK_pitch.pdf")
PREVIEW = os.path.join(ONE, "preview.png")
NMAP = os.path.join(ONE, "number_map.csv")

# ------------------------------------------------------------------ model numbers (data_only = Excel-cached values)
wb = load_workbook(XLSX, data_only=True)
nmap = []   # (label, value as shown, raw value, source)

def by_name(nm):
    dn = wb.defined_names[nm].attr_text
    sh, ref = dn.split("!")
    ref = ref.replace("$", "")
    v = wb[sh][ref].value
    if v is None:
        sys.exit(f"STOP: {nm} ({sh}!{ref}) has no cached value; rebuild + Excel save first")
    return float(v), f"{sh}!{ref} (defined name {nm})"

def by_label(sheet, label, col_header=None, col=None):
    """Locate a row by its column-A label; value from `col` or from the column whose header-row text equals col_header."""
    ws = wb[sheet]
    if col is None:
        hdr = None
        for r in range(1, 6):
            for c in range(1, ws.max_column + 1):
                if ws.cell(r, c).value == col_header:
                    hdr = c
        if hdr is None:
            sys.exit(f"STOP: header {col_header!r} not found on {sheet}")
        col = hdr
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            v = ws.cell(r, col).value
            if v is None:
                sys.exit(f"STOP: {sheet}!{ws.cell(r, col).coordinate} has no cached value")
            return float(v), f"{sheet}!{ws.cell(r, col).coordinate} (row label {label!r})"
    sys.exit(f"STOP: label {label!r} not found on {sheet}")

M = {}
for k in ("price", "price_perp", "price_exit", "upside_perp", "wacc_calc", "rf", "beta", "erp", "tg", "exit_mult",
          "shares", "cash", "px_comps_med", "deck_ev_ebitda", "comps_med", "g_fy27", "g_fy31", "ebit_margin", "mc_runs",
          "mc_guardrail_wacc_minus_g_min", "implied_margin", "g_fy28", "g_fy29", "g_fy30"):
    M[k] = by_name(k)
M["mc_p10"] = by_label("MC", "P10", col=2)
M["mc_p50"] = by_label("MC", "P50", col=2)
M["mc_p90"] = by_label("MC", "P90", col=2)
M["mc_min"] = by_label("MC", "min", col=2)
M["ebit_margin_fy26"] = by_label("Historical", "EBIT margin (income from operations / revenue)", col_header="FY2026")

# formatting helpers: whole dollars in text (price line keeps cents), percents to 1 decimal max
usd0 = lambda v: f"${v:,.0f}"
pct1 = lambda v: f"{v * 100:.1f}".rstrip("0").rstrip(".") + "%"
x1 = lambda v: f"{v:.1f}x"

def mnum(key, shown, label):
    raw, src = M[key]
    nmap.append((label, shown, raw, src))
    return shown

price_txt = mnum("price", f"${M['price'][0]:,.2f}", "current share price, Sep 21 2026")
dcf_txt = mnum("price_perp", usd0(M["price_perp"][0]), "DCF value per share, perpetuity method")
up_txt = mnum("upside_perp", f"+{pct1(M['upside_perp'][0])}", "upside, perpetuity DCF vs price")
exit_txt = mnum("price_exit", usd0(M["price_exit"][0]), "DCF value per share, exit multiple method")
evx_txt = mnum("deck_ev_ebitda", x1(M["deck_ev_ebitda"][0]), "DECK EV / LTM EBITDA at current price")
peer_txt = mnum("comps_med", x1(M["comps_med"][0]), "peer median EV / LTM EBITDA")
m26_txt = mnum("ebit_margin_fy26", pct1(M["ebit_margin_fy26"][0]), "FY26 EBIT margin (Historical tab)")
cash_m = M["cash"][0] / 1000
cash_txt = mnum("cash", f"${cash_m:,.0f}M", "cash and cash equivalents, Jun 30 2026 ($M)")
comps_txt = mnum("px_comps_med", usd0(M["px_comps_med"][0]), "comps implied price at peer median multiple")
p10_txt = mnum("mc_p10", usd0(M["mc_p10"][0]), "Monte Carlo P10")
p50_txt = mnum("mc_p50", usd0(M["mc_p50"][0]), "Monte Carlo P50")
p90_txt = mnum("mc_p90", usd0(M["mc_p90"][0]), "Monte Carlo P90")
min_txt = mnum("mc_min", usd0(M["mc_min"][0]), "Monte Carlo worst run of 10,000")
wacc_txt = mnum("wacc_calc", pct1(M["wacc_calc"][0]), "WACC (live formula)")
rf_txt = mnum("rf", pct1(M["rf"][0]), "risk-free rate")
beta_txt = mnum("beta", f"{M['beta'][0]:.2f}", "beta, Damodaran shoe")
erp_txt = mnum("erp", pct1(M["erp"][0]), "equity risk premium")
tg_txt = mnum("tg", pct1(M["tg"][0]), "terminal growth")
exm_txt = mnum("exit_mult", x1(M["exit_mult"][0]), "exit multiple EV/EBITDA")
sh_txt = mnum("shares", f"{M['shares'][0] / 1e6:.1f}M", "diluted shares (count)")
g27_txt = mnum("g_fy27", pct1(M["g_fy27"][0]), "forecast revenue growth FY27")
g31_txt = mnum("g_fy31", pct1(M["g_fy31"][0]), "forecast revenue growth FY31")
m_txt = mnum("ebit_margin", pct1(M["ebit_margin"][0]), "forecast EBIT margin")
runs_txt = mnum("mc_runs", f"{M['mc_runs'][0]:,.0f}", "Monte Carlo runs")
guard_txt = mnum("mc_guardrail_wacc_minus_g_min", pct1(M["mc_guardrail_wacc_minus_g_min"][0]), "Monte Carlo floor on WACC minus g")
im_txt = mnum("implied_margin", pct1(M["implied_margin"][0]), "market-implied EBIT margin at current price (DCF tab)")
g28_txt = mnum("g_fy28", pct1(M["g_fy28"][0]), "forecast revenue growth FY28")
g29_txt = mnum("g_fy29", pct1(M["g_fy29"][0]), "forecast revenue growth FY29")
g30_txt = mnum("g_fy30", pct1(M["g_fy30"][0]), "forecast revenue growth FY30")
# multiples, ROIC and RONIC rows read by defined name; cash components parsed from the Assumptions source text
for k in ("deck_ev_ebitda_stockanalysis", "deck_ev_ebitda_gurufocus", "deck_ev_ebitda_10y_median_gurufocus", "roic_fy26", "ronic", "px_at_ronic", "ronic_implied"):
    M[k] = by_name(k)
sa_txt = mnum("deck_ev_ebitda_stockanalysis", x1(M["deck_ev_ebitda_stockanalysis"][0]), "DECK EV/EBITDA per stockanalysis.com, Sep 21 2026")
gf_txt = mnum("deck_ev_ebitda_gurufocus", x1(M["deck_ev_ebitda_gurufocus"][0]), "DECK EV/EBITDA per GuruFocus, Sep 1 2026")
med10_txt = mnum("deck_ev_ebitda_10y_median_gurufocus", x1(M["deck_ev_ebitda_10y_median_gurufocus"][0]), "DECK 10-year median EV/EBITDA per GuruFocus, Sep 1 2026")
roic_txt = mnum("roic_fy26", pct1(M["roic_fy26"][0]), "FY26 ROIC on equity less cash plus lease liabilities (Historical tab)")
ronic_txt = mnum("ronic", pct1(M["ronic"][0]), "RONIC sensitivity input (Assumptions)")
pxr_txt = mnum("px_at_ronic", usd0(M["px_at_ronic"][0]), "perpetuity value per share at the RONIC input (DCF tab)")
rimp_txt = mnum("ronic_implied", f"{M['ronic_implied'][0] * 100:.0f}%", "implied return on new capital (DCF tab)")
import re as _re
_wa = wb["Assumptions"]
_r11 = next(rr for rr in range(4, _wa.max_row + 1) if _wa.cell(rr, 1).value == "cash")
_txt11 = str(_wa.cell(_r11, 4).value)
_bs_cash = float(_re.search(r"balance sheet ([0-9,]+) at", _txt11).group(1).replace(",", ""))
_rep = float(_re.search(r"less ([0-9,]+) repurchased", _txt11).group(1).replace(",", ""))
bs_cash_txt, rep_txt = f"${_bs_cash / 1000:,.0f}M", f"${_rep / 1000:,.0f}M"
nmap.append(("cash on the Jun 30 2026 balance sheet ($M)", bs_cash_txt, _bs_cash, f"Assumptions!D{_r11} source text; 10-Q Q1 FY27 balance sheet"))
nmap.append(("cash used for repurchases Jul 1 to Jul 9 2026 ($M)", rep_txt, _rep, f"Assumptions!D{_r11} source text; 10-Q Q1 FY27 subsequent events"))

# filing facts: typed, with source
F = [
    ("FY27 net sales guide, low", "$5.86B", 5.86e9, "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md)"),
    ("FY27 net sales guide, high", "$5.91B", 5.91e9, "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md)"),
    ("FY28 to FY30 framework: sales growth", "high-single-digit", "", "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md)"),
    ("FY28 to FY30 framework: operating margin", "low 20s", "", "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md)"),
    ("HOKA Q1 FY27 net sales growth", "+7.7%", 0.077, "10-Q Q1 FY27"),
    ("HOKA Q1 FY27 DTC growth", "+17%", 0.173, "10-Q Q1 FY27, HOKA brand net sales by channel: Direct-to-Consumer +17.3% (filings/10Q_Q1FY27.md line 1348)"),
    ("FY26 free cash flow", "topped $1B", 1e9, "8-K May 21 2026, CFO statement: over one billion dollars of free cash flow (filings/8K_releases_May_Jul_2026.md line 120)"),
    ("FY26 shares repurchased", "10.5M", 10.5e6, "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md line 114)"),
    ("FY26 repurchase spend", "$1.075B", 1.075e9, "8-K May 21 2026 (filings/8K_releases_May_Jul_2026.md line 114)"),
    ("debt", "0", 0, "10-K FY2026; Assumptions!debt"),
    ("HOKA net sales growth FY25", "+23.6%", 0.236, "10-K FY2026 segment note, Reportable Operating Segments (filings/10K_FY2026_Item7_FS.md: HOKA net sales 2,233,090 FY25 vs 1,806,740 FY24)"),
    ("HOKA net sales growth FY26", "+15.9%", 0.159, "10-K FY2026 segment note, Reportable Operating Segments (filings/10K_FY2026_Item7_FS.md: HOKA net sales 2,587,330 FY26 vs 2,233,090 FY25)"),
    ("go-forward tariff rate", "12.5%", 0.125, "Q1 FY27 earnings call, Jul 23 2026, per SGI Europe coverage"),
    ("prior tariff rate", "10%", 0.10, "Q1 FY27 earnings call, Jul 23 2026, per SGI Europe coverage"),
    ("buyback sizing, share of FY27 FCF", "80%", 0.80, "8-K Jul 23 2026 (filings/8K_releases_May_Jul_2026.md)"),
    ("Q2 FY27 catalyst date context", "Q2 FY27", "", "company reporting calendar"),
    ("cash date", "Jun 30 2026", "", "10-Q Q1 FY27 balance sheet"),
]
for lab, shown, raw, src in F:
    nmap.append((lab, shown, raw, src))



# ------------------------------------------------------------------ financial snapshot (Historical actuals, DCF forecasts)
def cell_by(sheet, label, col_header, hdr_rows=range(1, 4)):
    ws = wb[sheet]
    col = None
    for r in hdr_rows:
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value == col_header:
                col = c
    if col is None:
        sys.exit(f"STOP: header {col_header!r} not found on {sheet}")
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label:
            v = ws.cell(r, col).value
            if v is None:
                sys.exit(f"STOP: {sheet}!{ws.cell(r, col).coordinate} is empty")
            return float(v), f"{sheet}!{ws.cell(r, col).coordinate}"
    sys.exit(f"STOP: label {label!r} not found on {sheet}")

DCF_LAB = {}
for r in range(1, wb["DCF"].max_row + 1):
    v = wb["DCF"].cell(r, 1).value
    if isinstance(v, str):
        for key, start in (("rev", "Revenue"), ("g", "Revenue growth"), ("m", "EBIT margin"), ("ebitda", "EBITDA"), ("fcff", "FCFF")):
            if v.startswith(start) and key not in DCF_LAB and not (key == "rev" and v.startswith("Revenue growth")):
                DCF_LAB[key] = v
SNAP_COLS = ["FY2024A", "FY2025A", "FY2026A", "FY2027E", "FY2031E"]
snap = {k: {} for k in ("rev", "g", "m", "ebitda", "fcff")}
for c in SNAP_COLS:
    fy = c[:-1]
    if c.endswith("A"):
        rev, s_rev = cell_by("Historical", "Net sales", fy)
        g, s_g = cell_by("Historical", "Revenue growth", fy)
        m, s_m = cell_by("Historical", "EBIT margin (income from operations / revenue)", fy)
        ebit, s_ebit = cell_by("Historical", "Income from operations", fy)
        da, s_da = cell_by("Historical", "Depreciation amortization and accretion", fy)
        snap["rev"][c] = (rev, s_rev); snap["g"][c] = (g, s_g); snap["m"][c] = (m, s_m)
        snap["ebitda"][c] = (ebit + da, f"{s_ebit} + {s_da} (income from operations + D&A and accretion)")
    else:
        for key in ("rev", "g", "m", "ebitda", "fcff"):
            snap[key][c] = cell_by("DCF", DCF_LAB[key], c, hdr_rows=range(1, 3))
SNAP_ROWS = [("rev", "Net sales", lambda v: f"{v / 1000:,.0f}"), ("g", "Growth %", lambda v: f"{v * 100:.1f}%"),
             ("m", "EBIT margin %", lambda v: f"{v * 100:.1f}%"), ("ebitda", "EBITDA", lambda v: f"{v / 1000:,.0f}"),
             ("fcff", "FCFF", lambda v: f"{v / 1000:,.0f}")]
snap_html = []
for key, lab, fmt in SNAP_ROWS:
    cells = []
    for c in SNAP_COLS:
        if c in snap[key]:
            raw, src = snap[key][c]
            shown = fmt(raw)
            nmap.append((f"snapshot {lab} {c}", shown, raw, src))
            cells.append(f"<td>{shown}</td>")
        else:
            cells.append("<td></td>")
    snap_html.append(f"<tr><th>{lab}</th>{''.join(cells)}</tr>")
snap_table = ('<h2 style="margin-top:9pt">Financial snapshot, FY2024 to FY2031 ($M unless noted)</h2><table class="snap"><thead><tr><th></th>'
              + "".join(f"<th>{c}</th>" for c in SNAP_COLS) + "</tr></thead><tbody>" + "".join(snap_html) + "</tbody></table>"
              '<div class="cap" style="margin-top:2pt">Actuals come from the Historical tab, sourced to the 10-K FY2026 and 10-K FY2024, and forecast years come from the DCF tab. FCFF is a forecast-only line.</div>')

# ------------------------------------------------------------------ HTML
football = os.path.join(ONE, "football.png").replace("\\", "/")
hist = os.path.join(ROOT, "mc", "hist.png").replace("\\", "/")
html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>DECK pitch</title>
<style>
@page {{ size: Letter portrait; margin: 0.38in 0.45in 0.34in 0.45in; }}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; }}
body {{ font-family: "Segoe UI", Arial, Helvetica, sans-serif; font-size: 9.1pt; line-height: 1.25; color: #0b0b0b; hyphens: none; -webkit-hyphens: none; overflow-wrap: normal; word-break: normal; }}
.nb {{ white-space: nowrap; }}
.page {{ height: 10.24in; display: flex; flex-direction: column; }}
.banner {{ background: #1F3864; color: #fff; padding: 6pt 10pt; font-size: 10.8pt; font-weight: 600; letter-spacing: 0; white-space: nowrap; overflow: hidden; }}
.banner .buy {{ background: #eb6834; color: #fff; padding: 0 5pt; border-radius: 2pt; margin: 0 2pt; }}
.sub {{ color: #52514e; font-size: 10pt; font-style: italic; margin: 4pt 0 6pt 0; }}
.cols {{ display: flex; gap: 14pt; }}
.left {{ flex: 0 0 47%; }}
.right {{ flex: 1 1 auto; }}
h2 {{ font-size: 9.2pt; letter-spacing: 1pt; color: #1F3864; border-bottom: 1.2pt solid #1F3864; margin: 7pt 0 3pt 0; padding-bottom: 1pt; text-transform: uppercase; }}
h2:first-child {{ margin-top: 0; }}
ol, ul {{ margin: 0; padding-left: 13pt; }}
li {{ margin: 0 0 3pt 0; }}
p {{ margin: 0 0 3pt 0; }}
b.k {{ color: #1F3864; }}
.num {{ color: #2a78d6; font-weight: 600; }}
.right img {{ width: 100%; display: block; border: 0.5pt solid #e1e0d9; }}
.cap {{ font-size: 8.2pt; color: #52514e; margin: 2pt 0 6pt 0; }}
.val {{ background: #f3f5f9; border-left: 2.5pt solid #2a78d6; padding: 4pt 8pt; }}
table.snap {{ border-collapse: collapse; width: 100%; font-size: 8.8pt; margin-top: 2pt; }}
table.snap th, table.snap td {{ padding: 1.5pt 6pt; text-align: right; border-bottom: 0.5pt solid #e1e0d9; }}
table.snap th:first-child, table.snap td:first-child {{ text-align: left; }}
table.snap thead th {{ color: #1F3864; border-bottom: 1pt solid #1F3864; font-weight: 600; }}
table.snap tbody th {{ font-weight: 400; color: #0b0b0b; }}
.foot {{ border-top: 0.8pt solid #898781; margin-top: auto; padding-top: 4pt; font-size: 7.8pt; color: #52514e; line-height: 1.25; }}
</style></head><body><div class="page">
<div class="banner">Deckers Outdoor (NYSE: DECK) &nbsp;|&nbsp; <span class="buy">BUY</span> &nbsp;|&nbsp; DCF value {dcf_txt} &nbsp;|&nbsp; Price {price_txt} (Sep 21 2026) &nbsp;|&nbsp; Upside {up_txt}</div>
<div class="sub">DECK trades at {evx_txt} EV/LTM EBITDA with zero debt. Management's FY27 guide is net sales of $5.86B to $5.91B.</div>
<div class="cols">
<div class="left">
<h2>Thesis at {price_txt}</h2>
<ol>
<li><b class="k">Cheapest multiple in the peer set.</b> On stockanalysis numbers DECK trades at <span class="num">{sa_txt}</span> EV/EBITDA against an <span class="num">{peer_txt}</span> peer median. On GuruFocus numbers it trades at <span class="num">{gf_txt}</span> against its own {med10_txt} ten-year median. The FY28 to FY30 framework calls for high-single-digit sales growth and operating margin in the low 20s.</li>
<li><b class="k">FY27 at guidance, out-years below the framework.</b> FY27 growth of {g27_txt} equals the guide midpoint. FY29 and FY30 growth of {g29_txt} and {g30_txt} runs under the high-single-digit framework. Margin holds at {m_txt} against a Jul 23 guide of slightly better.</li>
<li><b class="k">Cash generation funds the buyback.</b> FY26 free cash flow topped $1B and DECK retired 10.5M shares for $1.075B (8-K May 21 2026). HOKA grew +7.7% in Q1 FY27 with DTC +17% (10-Q). Zero debt.</li>
</ol>
<h2>Valuation: four methods</h2>
<div class="val">Four methods land above the price. The perpetuity DCF gives <span class="num">{dcf_txt}</span> and the {exm_txt} exit multiple gives <span class="num">{exit_txt}</span>. Comps at the peer median give <span class="num">{comps_txt}</span>. The Monte Carlo P10 to P90 band runs <span class="num">{p10_txt} to {p90_txt}</span>. The price implies a <span class="num">{im_txt}</span> EBIT margin, held flat from FY27, against the {m_txt} forecast. The perpetuity value implies a <span class="num">{rimp_txt}</span> return on new capital, close to FY26's <span class="num">{roic_txt}</span> ROIC. At {ronic_txt} it is <span class="num">{pxr_txt}</span>.</div>
<h2>Risks: HOKA growth and tariffs</h2>
<ol>
<li><b class="k">HOKA growth is decelerating.</b> HOKA net sales grew +23.6% in FY25 and +15.9% in FY26, per the segment note in the FY2026 10-K. The Q1 FY27 10-Q shows +7.7%.</li>
<li><b class="k">Tariff rate went from 10% to 12.5%.</b> On the Jul 23 2026 Q1 FY27 call, per SGI Europe coverage, the go-forward tariff rate moved to 12.5% from 10%.</li>
</ol>
<h2>Catalysts: Q2 FY27 and holiday quarter</h2>
<p>The Q2 FY27 report will show whether HOKA DTC growth holds after the +17% Q1 print (10-Q). Then the holiday quarter, UGG's seasonal peak. The EPS guide assumes buybacks of about 80% of FY27 free cash flow.</p>
</div>
<div class="right">
<img src="file:///{football}" alt="football field">
<div class="cap">Each bar runs low to high. The dotted line marks the current price and the diamond the {dcf_txt} base DCF, while the analyst-target and 52-week bars come from stockanalysis.com and the other four from the model.</div>
<img src="file:///{hist}" alt="Monte Carlo histogram">
<div class="cap">{runs_txt} independent draws over growth, margin, WACC and terminal g, each repriced through the perpetuity DCF with the WACC minus g floor of {guard_txt} enforced. P50 is {p50_txt}. Growth and margin are drawn independently but move together through operating leverage, so the true spread is modestly wider.</div>
</div>
</div>
{snap_table}
<div class="foot">WACC {wacc_txt} is Rf {rf_txt} plus beta {beta_txt} (Damodaran shoe) times ERP {erp_txt}, with terminal g {tg_txt} and a {exm_txt} EV/EBITDA exit multiple as the cross-check. {sh_txt} diluted shares. Cash {cash_txt} (Jun 30 {bs_cash_txt} less {rep_txt} repurchased to Jul 9), debt 0. Leases stay in EBIT and SBC is expensed. Sources: 10-K FY2026, 10-K FY2024, 10-Q Q1 FY2027, 8-K releases May 21 and Jul 23 2026, SGI Europe call coverage, stockanalysis.com, GuruFocus, TipRanks, Damodaran, FRED. Salman Koshy, UC Berkeley. Independent student project, not investment advice.</div>
</div></body></html>
"""
def nowrap_compounds(doc):
    """Wrap hyphenated compounds (high-single-digit, 10-K, 52-week ...) in nowrap spans so no compound splits at a line end."""
    head, body_ = doc.split("<body>", 1)
    out = []
    for piece in re.split(r"(<[^>]+>)", body_):
        out.append(piece if piece.startswith("<") else re.sub(r"\b(\w+(?:-\w+)+)\b", r'<span class="nb">\1</span>', piece))
    return head + "<body>" + "".join(out)
import re
html = nowrap_compounds(html)
HYPHENATED = sorted(set(re.findall(r"\b\w+(?:-\w+)+\b", re.sub(r"<[^>]+>", " ", html.split("<body>", 1)[1]))))
assert "\u2014" not in html and "\u2013" not in html, "em/en dash in page"
import re as _re
_txt = _re.sub(r"<[^>]+>", " ", html.split("<body>", 1)[1])
_odd = sorted({ch for ch in _txt if ord(ch) > 126 and ch not in "\u00a0"})
print("non-ASCII characters in page text:", _odd if _odd else "none")
open(HTML, "w", encoding="utf-8", newline="").write(html)

# ------------------------------------------------------------------ number map
with open(NMAP, "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh, lineterminator="\n")
    w.writerow(["label", "shown_on_page", "raw_value", "source"])
    w.writerows(nmap)

# ------------------------------------------------------------------ PDF via headless Edge / Chrome, playwright fallback
def print_pdf():
    cands = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
             r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"]
    for exe in cands:
        if os.path.exists(exe):
            if os.path.exists(PDF):
                os.remove(PDF)
            cmd = [exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--no-first-run", "--disable-extensions",
                   f"--print-to-pdf={PDF}", "--virtual-time-budget=4000", "file:///" + HTML.replace("\\", "/")]
            subprocess.run(cmd, capture_output=True, timeout=120)
            if os.path.exists(PDF) and os.path.getsize(PDF) > 1000:
                return os.path.basename(exe)
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch(); pg = b.new_page(); pg.goto("file:///" + HTML.replace("\\", "/"))
            pg.pdf(path=PDF, format="Letter", prefer_css_page_size=True, print_background=True); b.close()
        return "playwright"
    except Exception as ex:
        sys.exit(f"STOP: no headless browser produced the PDF ({ex!r})")

engine = print_pdf()

# ------------------------------------------------------------------ page count, overflow, preview
from pypdf import PdfReader
n_pages = len(PdfReader(PDF).pages)
import pypdfium2 as pdfium
doc = pdfium.PdfDocument(PDF)
page = doc[0]
w_in, h_in = page.get_width() / 72, page.get_height() / 72
img = page.render(scale=200 / 72).to_pil()
img.save(PREVIEW)
print(f"engine: {engine}")
print(f"PDF: {PDF}")
print(f"page count: {n_pages}   page size: {w_in:.2f} x {h_in:.2f} in ({'Letter portrait' if abs(w_in - 8.5) < 0.01 and abs(h_in - 11) < 0.01 else 'NOT LETTER'})")
print("OVERFLOW WARNING: content spilled onto page 2 or more" if n_pages != 1 else "overflow: none (exactly one page)")
print(f"preview: {PREVIEW} ({img.width} x {img.height} px at 200 dpi)")
layer = page.get_textpage().get_text_range()
flat = re.sub(r"\s+", " ", layer)
missing = [h for h in HYPHENATED if h not in flat]
broken = re.findall(r"\w+-?[\ufffe\u00ad]\w+|\w+-\r?\n\w+", layer)
glued = [h for h in HYPHENATED if h.replace("-", "") in flat]
print(f"hyphen check: {len(HYPHENATED)} hyphenated compounds in source: {HYPHENATED}")
print(f"hyphen check: missing from text layer {missing or 'none'}; split at a line end {broken or 'none'}; glued forms {glued or 'none'}")
print("hyphen check:", "PASS" if not (missing or broken or glued) else "FAIL")
print(f"number_map.csv: {NMAP} ({len(nmap)} rows)")
with open(NMAP, encoding="utf-8") as fh:
    print(fh.read())
